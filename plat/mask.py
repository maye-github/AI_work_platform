# -*- coding: utf-8 -*-
"""脱敏 —— 明文 ↔ 代号 の対照表。

**なぜ「置換スクリプト」ではなく「対照表」か**
対照表（`registry/MASK_MAP.yaml`）が唯一の写像であり、成果物には**代号しか書かない**。
こうすると「どの成果物に何が漏れているか」を機械で判定でき、対外版の作り直しも 1 コマンドで済む。
対照表そのものが機密なので `.gitignore` 済み（コミット禁止）。

使い方:
    plat mask assign --plain "<顧客名>" --code "<客户A>"     # 対照を登録
    plat mask apply  [--write]                              # 成果物の明文を代号へ置換
    plat mask check                                         # 残留明文 / 未登録代号を検査（gate6 の一部）
    plat mask list [--safe]                                 # 対照一覧（--safe は明文を伏せる）
"""
from __future__ import annotations

import re
from pathlib import Path

import yaml

from . import model, paths, rules_data as R


def load_map() -> list[dict]:
    p = paths.mask_map_file()
    if not p.is_file():
        return []
    try:
        d = model.load_yaml_doc(p)
    except Exception:
        return []
    return list(d.get("entries") or [])


def save_map(entries: list[dict]) -> None:
    p = paths.mask_map_file()
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_text(
        yaml.dump({"entries": entries}, allow_unicode=True, sort_keys=False,
                  default_flow_style=False, width=4096),
        encoding="utf-8",
    )


def assign(plain: str, code: str, category: str = "customer") -> dict:
    entries = load_map()
    for e in entries:
        if e.get("plain") == plain:
            raise SystemExit(f"既に登録済みです: {plain} → {e.get('code')}")
        if e.get("code") == code:
            raise SystemExit(f"代号が重複しています: {code}")
    if not (code.startswith("<") and code.endswith(">")):
        code = f"<{code}>"
    e = {"code": code, "plain": plain, "category": category}
    entries.append(e)
    save_map(entries)
    return e


def _artifact_texts() -> list[Path]:
    """検査対象は成果物のみ。`raw/`（素材）と生成物は対象外。

    `raw/` は原文のままであり、`<host>:<port>` のような**書き方のための山括弧**を
    正当に含む。ここを検査対象にすると誤検出で門番が使えなくなる。
    """
    out: list[Path] = []
    for _case, _module, mod in paths.modules():
        for p in mod.rglob("*"):
            if not p.is_file() or p.suffix.lower() not in (".yaml", ".yml", ".md"):
                continue
            parts = p.relative_to(mod).parts
            if any(x in parts for x in ("raw", "ai_input", "readable",
                                        "node_modules", "dist", ".venv", "venv",
                                        "__pycache__", ".git")):
                continue
            if p.name in ("module.yaml", "SESSION_HANDOFF.md"):
                continue
            out.append(p)
    return out


def apply(write: bool = False) -> list[tuple[Path, int]]:
    entries = load_map()
    changed: list[tuple[Path, int]] = []
    for p in _artifact_texts():
        text = p.read_text(encoding="utf-8")
        new, n = text, 0
        for e in entries:
            plain = str(e.get("plain") or "")
            if not plain:
                continue
            n += new.count(plain)
            new = new.replace(plain, str(e.get("code")))
        if n:
            changed.append((p, n))
            if write:
                p.write_text(new, encoding="utf-8")
    return changed


# 代号の形：`<客户A>` のような短い山括弧。HTML タグと紛れないよう長さと字種を絞る。
_CODE_LIKE = re.compile(r"<([A-Za-z0-9_぀-ヿ一-鿿]{1,10})>")
_HTML_TAGS = {
    "br", "b", "i", "u", "p", "em", "strong", "code", "pre", "ul", "ol", "li",
    "table", "tr", "td", "th", "h1", "h2", "h3", "h4", "div", "span", "a", "hr",
}
# 書き方のための山括弧（プレースホルダ）。代号ではないので指摘しない。
_TEMPLATE_PLACEHOLDERS = {
    "host", "port", "scheme", "path", "name", "value", "id", "n", "x",
    "案件", "模块", "人", "条目", "成果物ID", "案件名", "模块名", "例", "…",
}

# `plat mask prep` が形態で置換した自動代号（<IP_1> / <MAIL_1> 等）。
_AUTO_PLACEHOLDER = re.compile(
    r"^(?:IP|CONN|MAIL|KEY|TOKEN|PWD|FNC|TRG|EMP|NAME|VAL|COL|TBL|"
    r"BUCKET|DOM|PROFILE|SID|PORT)_\d+$")


def check() -> list[str]:
    """成果物に残る明文 / 未登録の代号を返す。空リスト = きれい。"""
    err: list[str] = []
    entries = load_map()
    codes = {str(e.get("code")) for e in entries}
    for p in _artifact_texts():
        text = p.read_text(encoding="utf-8")
        for e in entries:
            plain = str(e.get("plain") or "")
            if plain and plain in text:
                err.append(f"{paths.rel(p)}: 明文が残っています（{e.get('code')}）")
        for m in _CODE_LIKE.finditer(text):
            tok = m.group(0)
            inner = m.group(1)
            if (tok in codes or inner.lower() in _HTML_TAGS
                    or inner.lower() in _TEMPLATE_PLACEHOLDERS
                    or _AUTO_PLACEHOLDER.match(inner)):
                continue
            err.append(f"{paths.rel(p)}: 未登録の代号 {tok}")
    return err


# =============================================================================
# 取込前の脱敏（raw/ → ai_input/）
# =============================================================================
_PATTERN_PREFIX = {
    "内部IP(RFC1918)": "IP",
    "接続文字列(ユーザ:パス)": "CONN",
    "社内メール": "MAIL",
    "メール(取込)": "MAIL",
    "AWSアクセスキー": "KEY",
    "秘密鍵ブロック": "KEY",
    "Bearer トークン": "TOKEN",
    "パスワード代入": "PWD",
    "復号関数の実引数": "FNC",
    "緊急配信トリガー式": "TRG",
}

# 取込脱敏は「漏らさない」側に倒す。G6 のメール regex は test. を除外するが、
# 取込では example.com/org/jp 以外は実メールとして扱う。
_PREP_PATTERNS = list(R.SECRET_PATTERNS) + [
    ("メール(取込)",
     r"\b[\w.+-]+@(?!example\.(?:com|org|jp)\b)[\w-]+(?:\.[\w-]+)+\b"),
]

_CANDIDATE = re.compile(
    r"(?:株式会社|有限会社|様|氏名|担当|顧客|プロジェクト|会社名|社員コード|事業場)")


def _rel_or_str(p: Path) -> str:
    try:
        return paths.rel(p)
    except Exception:
        return str(p)


def _extract_text(path: Path) -> str | None:
    """テキスト化できる素材だけを読む。PDF/画像は未対応として返す。"""
    suffix = path.suffix.lower()
    if suffix in (".md", ".txt", ".csv", ".tsv", ".json", ".yaml", ".yml",
                  ".log", ".sql", ".html", ".htm"):
        return path.read_text(encoding="utf-8", errors="replace")
    if suffix == ".docx":
        import docx
        doc = docx.Document(str(path))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append("\t".join(c.text for c in row.cells))
        return "\n".join(parts)
    if suffix in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        parts: list[str] = []
        for ws in wb.worksheets:
            parts.append(f"## sheet: {ws.title}")
            for row in ws.iter_rows(values_only=True):
                parts.append("\t".join("" if v is None else str(v) for v in row))
        return "\n".join(parts)
    return None


def _mask_text(text: str, entries: list[dict], stats: dict) -> str:
    # 1) 登録済みの明文（長いものから先に置換して部分一致を避ける）
    for e in sorted(entries, key=lambda x: len(str(x.get("plain") or "")), reverse=True):
        plain = str(e.get("plain") or "")
        if not plain:
            continue
        n = text.count(plain)
        if n:
            text = text.replace(plain, str(e.get("code")))
            stats["known"] += n
    # 2) 形態で検出できる L3
    for name, rx in _PREP_PATTERNS:
        prefix = _PATTERN_PREFIX.get(name, "SEC")
        counter = {"n": stats["patterns"].get(prefix, 0)}

        def repl(_m, prefix=prefix, counter=counter):
            counter["n"] += 1
            return f"<{prefix}_{counter['n']}>"

        text = re.sub(rx, repl, text)
        stats["patterns"][prefix] = counter["n"]
    return text


def _default_src(case: str, module: str) -> Path:
    mod = paths.module_dir(case, module)
    for rel in ("requirements/raw", "raw"):
        p = mod / rel
        if p.is_dir():
            return p
    return mod / "requirements" / "raw"


def prep(case: str = "", module: str = "", src: str | Path | None = None,
         out: str | Path | None = None) -> dict:
    """raw/ を読み、脱敏コピーを ai_input/ に作る。**AI は raw/ を読まない。**"""
    entries = load_map()
    src_p = Path(src) if src else _default_src(case, module)
    out_p = Path(out) if out else paths.module_dir(case, module) / "requirements" / "ai_input"
    if not src_p.is_dir():
        raise SystemExit(f"raw がありません: {src_p}")
    out_p.mkdir(parents=True, exist_ok=True)

    stats = {"known": 0, "patterns": {}, "candidates": 0}
    files: list[str] = []
    unsupported: list[str] = []

    for p in sorted(src_p.rglob("*")):
        if not p.is_file():
            continue
        text = _extract_text(p)
        if text is None:
            unsupported.append(p.relative_to(src_p).as_posix())
            continue
        stats["candidates"] += len(_CANDIDATE.findall(text))
        masked = _mask_text(text, entries, stats)
        rel = p.relative_to(src_p).with_suffix(".md")
        dst = out_p / rel
        dst.parent.mkdir(parents=True, exist_ok=True)
        header = (
            f"# 取込（脱敏済み）: {p.name}\n\n"
            f"> 出典: {_rel_or_str(p)}\n"
            f"> 生成: plat mask prep（原文は AI に渡さない）\n\n"
        )
        dst.write_text(header + masked, encoding="utf-8", newline="\n")
        files.append(rel.as_posix())

    pattern_parts = ", ".join(f"{k}={v}" for k, v in sorted(stats["patterns"].items())) or "なし"
    report = [
        "# MASK_REPORT",
        "",
        f"- 原文: `{_rel_or_str(src_p)}`（AI には渡さない）",
        f"- 脱敏後: `{_rel_or_str(out_p)}`",
        f"- 変換ファイル: {len(files)} 件",
        f"- 既知の置換: {stats['known']} 件",
        f"- 形態で置換: {pattern_parts}",
        f"- 人名/公司名の候補: {stats['candidates']} 箇所（**未自動置換**。`MASK_MAP` に登録して再実行）",
        f"- 未対応ファイル: {', '.join(unsupported) if unsupported else 'なし'}",
        "",
        "> このレポートには明文を書かない。原文は `raw/` にのみ残す。",
    ]
    (out_p / "MASK_REPORT.md").write_text(
        "\n".join(report) + "\n", encoding="utf-8", newline="\n")

    return {"files": len(files), "known": stats["known"], "patterns": stats["patterns"],
            "unsupported": unsupported, "candidates": stats["candidates"], "out": str(out_p)}
