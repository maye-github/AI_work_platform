# -*- coding: utf-8 -*-
"""華語圏向けではない —— 対日成果物（Excel／Word）への書き出し。

**二重の形態（規格書 §9.3）**：機読源（YAML/MD）から人向けの正式版を自動生成する。
生成物は `export/`（gitignore 済み）に置き、**手で編集しない**（次回生成で消える）。

依存：`openpyxl`（必須）／`python-docx`（任意。無ければ Word は飛ばして報告する）
"""
from __future__ import annotations

import datetime as dt
from pathlib import Path

from . import model, paths

try:
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    HAS_XLSX = True
except Exception:  # pragma: no cover
    HAS_XLSX = False

try:
    import docx
    HAS_DOCX = True
except Exception:  # pragma: no cover
    HAS_DOCX = False

HEAD_FILL = "99CCFF"   # 参考書のヘッダ色に合わせる（gen_formal と同じ）
SUB_FILL = "CCECFF"


def _borders():
    thin = Side(style="thin", color="999999")
    return Border(left=thin, right=thin, top=thin, bottom=thin)


def _front_rows(a: model.Artifact) -> list[tuple[str, str]]:
    rows = []
    for k, v in a.data.items():
        if isinstance(v, list):
            v = ", ".join(str(x) for x in v)
        elif isinstance(v, dict):
            continue
        rows.append((str(k), "" if v is None else str(v)))
    return rows


_UNSAFE = '\\/:*?"<>|（）()　 '


def _safe_name(case: str, module: str, ident: str, title) -> str:
    """ファイル名に使えない文字を落とし、長すぎない名前にする（括弧で切れないように）。"""
    t = str(title or "").strip()
    for ch in _UNSAFE:
        t = t.replace(ch, "_")
    t = "".join(ch for ch in t if ch.isprintable() and ch not in _UNSAFE)
    t = t[:24].strip("_")
    base = f"{case}_{module}_{ident}"
    return f"{base}_{t}" if t else base


def write_xlsx(a: model.Artifact, out_dir: Path) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "概要"
    bd = _borders()
    ws["A1"] = a.data.get("title") or a.id
    ws["A1"].font = Font(bold=True, size=14)
    ws.merge_cells("A1:B1")
    r = 3
    for k, v in _front_rows(a):
        ws.cell(r, 1, k).font = Font(bold=True)
        ws.cell(r, 1).fill = PatternFill("solid", fgColor=SUB_FILL)
        ws.cell(r, 1).border = bd
        ws.cell(r, 2, v).border = bd
        ws.cell(r, 2).alignment = Alignment(wrap_text=True, vertical="top")
        r += 1
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 80

    # 本文の「dict のリスト」を 1 シートずつ
    for key, val in a.data.items():
        if not (isinstance(val, list) and val and all(isinstance(x, dict) for x in val)):
            continue
        cols: list[str] = []
        for item in val:
            for k2 in item:
                if k2 not in cols:
                    cols.append(str(k2))
        sh = wb.create_sheet(str(key)[:28])
        for j, c in enumerate(cols, start=1):
            cell = sh.cell(1, j, c)
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor=HEAD_FILL)
            cell.border = bd
            sh.column_dimensions[get_column_letter(j)].width = 24
        for i, item in enumerate(val, start=2):
            for j, c in enumerate(cols, start=1):
                v = item.get(c, "")
                if isinstance(v, (list, dict)):
                    v = str(v)
                sh.cell(i, j, v).border = bd
                sh.cell(i, j).alignment = Alignment(wrap_text=True, vertical="top")
        sh.freeze_panes = "A2"

    out_dir.mkdir(parents=True, exist_ok=True)
    out = out_dir / (_safe_name(a.case, a.module, a.id, a.data.get("title")) + ".xlsx")
    wb.save(out)
    return out


def write_docx(a: model.Artifact, out_dir: Path) -> Path:
    d = docx.Document()
    d.add_heading(str(a.data.get("title") or a.id), level=1)
    t = d.add_table(rows=0, cols=2)
    t.style = "Table Grid"
    for k, v in _front_rows(a):
        row = t.add_row().cells
        row[0].text = k
        row[1].text = v
    if a.body.strip():
        d.add_heading("本文", level=2)
        for para in a.body.split("\n\n"):
            if para.strip():
                d.add_paragraph(para.strip())
    out_dir.mkdir(parents=True, exist_ok=True)
    out = out_dir / f"{a.case}_{a.module}_{a.id}.docx"
    d.save(out)
    return out


def run(module: str | None = None, case: str | None = None,
        out_dir: Path | None = None, fmt: str = "xlsx") -> list[str]:
    if not HAS_XLSX and fmt == "xlsx":
        raise SystemExit("openpyxl がありません: pip install openpyxl")
    out_dir = out_dir or (paths.root() / "export")
    arts = [a for a in model.load_all() if a.id]
    if module:
        arts = [a for a in arts if a.module == module]
    if case:
        arts = [a for a in arts if a.case == case]
    if not arts:
        raise SystemExit("書き出す成果物がありません")
    log: list[str] = []
    skipped = 0
    for a in arts:
        if fmt == "xlsx":
            log.append(str(write_xlsx(a, out_dir)))
        elif fmt == "docx":
            if not HAS_DOCX:
                skipped += 1
                continue
            log.append(str(write_docx(a, out_dir)))
        else:
            raise SystemExit(f"未知の形式: {fmt}（xlsx / docx）")
    if skipped:
        log.append(f"（Word {skipped} 件は python-docx 未導入のため飛ばしました: "
                   "pip install python-docx）")
    return log
